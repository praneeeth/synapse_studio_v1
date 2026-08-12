"""Upload text extraction.

Previously inlined in the upload route and limited to PDF + UTF-8 text. Split
out here and extended with DOCX and a proper encoding fallback chain.
"""

from __future__ import annotations

import logging
from io import BytesIO
from typing import Tuple

logger = logging.getLogger(__name__)

TEXT_EXTENSIONS = {".txt", ".md", ".markdown", ".csv", ".json", ".log", ".rst", ".yaml", ".yml"}


class UnsupportedFileError(Exception):
    """Raised when no extractor can handle the upload."""


def _decode(content: bytes) -> str:
    for encoding in ("utf-8", "utf-8-sig", "utf-16", "cp1252"):
        try:
            return content.decode(encoding)
        except (UnicodeDecodeError, LookupError):
            continue
    return content.decode("latin-1", errors="ignore")


def _looks_binary(text: str) -> bool:
    """Detect content that decoded "successfully" but isn't really text.

    NUL and other control bytes survive str.strip(), so a truthiness check alone
    lets arbitrary binary blobs through as documents.
    """
    if not text:
        return True
    if "\x00" in text:
        return True

    printable = sum(1 for ch in text if ch.isprintable() or ch in "\n\r\t")
    return printable / len(text) < 0.85


def _extract_pdf(content: bytes) -> str:
    from pypdf import PdfReader

    reader = PdfReader(BytesIO(content))
    if reader.is_encrypted:
        try:
            reader.decrypt("")  # many PDFs are "encrypted" with an empty password
        except Exception as exc:  # noqa: BLE001
            raise UnsupportedFileError(f"PDF is password protected: {exc}") from exc

    pages = []
    for page_num, page in enumerate(reader.pages, start=1):
        try:
            pages.append(page.extract_text() or "")
        except Exception as exc:  # noqa: BLE001 - one bad page shouldn't kill the upload
            logger.warning("Failed to extract PDF page %d: %s", page_num, exc)
            pages.append("")
    return "\n\n".join(pages)


def _extract_docx(content: bytes) -> str:
    try:
        import docx  # python-docx
    except ImportError as exc:
        raise UnsupportedFileError("python-docx is not installed; cannot read .docx") from exc

    document = docx.Document(BytesIO(content))
    parts = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            parts.append("\t".join(cell.text for cell in row.cells))
    return "\n".join(parts)


def extract_text(content: bytes, filename: str, content_type: str = "") -> Tuple[str, str]:
    """Return (text, detected_kind) for an uploaded file.

    Raises UnsupportedFileError when the type cannot be handled or yields nothing.
    """
    name = (filename or "").lower()
    ctype = (content_type or "").lower()

    if ctype == "application/pdf" or name.endswith(".pdf"):
        text, kind = _extract_pdf(content), "pdf"
    elif name.endswith((".docx",)) or "wordprocessingml" in ctype:
        text, kind = _extract_docx(content), "docx"
    elif name.endswith(".doc"):
        raise UnsupportedFileError("Legacy .doc is not supported -- please convert to .docx or PDF.")
    elif any(name.endswith(ext) for ext in TEXT_EXTENSIONS) or ctype.startswith("text/"):
        text, kind = _decode(content), "text"
        if _looks_binary(text):
            raise UnsupportedFileError(
                f"{filename!r} has a text extension but contains binary data."
            )
    else:
        # Unknown extension: try decoding as text before giving up.
        text = _decode(content)
        if _looks_binary(text):
            raise UnsupportedFileError(
                f"Unsupported file type: {filename!r}. "
                "Supported: PDF, DOCX, and plain-text formats (txt, md, csv, json, ...)."
            )
        kind = "text"

    if not text.strip():
        raise UnsupportedFileError(
            "No readable text could be extracted. "
            "If this is a scanned PDF it needs OCR, which this build does not perform."
        )

    return text, kind
